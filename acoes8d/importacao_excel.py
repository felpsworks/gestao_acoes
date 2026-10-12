"""Leitura das planilhas A3 antigas (FM-QS-025, layout da rev. 12) e gravação no
módulo. Usado pelo botão "Importar A3 (Excel)" do Formulário A3 e pelo script de
linha de comando importar_a3_excel.py.

Regras:
  - Um A3 cujo ID já existe no sistema é IGNORADO (nunca sobrescreve o que já foi
    preenchido). Pode importar de novo sem duplicar.
  - Status das ações (coluna Status do FRP - A3): G -> Concluída (data de
    conclusão = prazo, a melhor informação disponível na planilha), Y/R -> Em
    andamento, vazio -> Aberta. Ações sem responsável (ou "NA") não são importadas.
  - Status do A3: mesma regra do sistema (ações G + etapas do 8D realizadas +
    "8D Aprovado" = Concluído).
"""
import json
import re
from datetime import date, datetime
from pathlib import Path

from .modelo_a3 import CATEGORIAS_6M, CHECKLIST, DIAS_PRAZO_LIMITE_A3, SETORES_PADRAO, dados_vazios, normalizar_setor  # noqa: F401
from .regras import status_automatico_a3

STATUS_FAROL = {"G": "Concluída", "Y": "Em andamento", "R": "Em andamento"}


def texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, (datetime, date)):
        return valor.strftime("%Y-%m-%d")
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def data_iso(valor):
    if isinstance(valor, (datetime, date)):
        return valor.strftime("%Y-%m-%d")
    t = texto(valor)
    m = re.match(r"(\d{1,2})/(\d{1,2})/(\d{2,4})", t)
    if m:
        d, mth, a = (int(x) for x in m.groups())
        return f"{a + 2000 if a < 100 else a:04d}-{mth:02d}-{d:02d}"
    return t[:10] if re.match(r"\d{4}-\d{2}-\d{2}", t) else ""


def pessoa_valida(nome):
    return bool(nome) and nome.strip().upper() not in {"NA", "N/A", "-", "—"}


def aba(wb, *nomes):
    for nome in nomes:
        for ws in wb.worksheets:
            if ws.title.strip().lower() == nome.lower():
                return ws
    return None


# Palavra-chave do rótulo da linha no Excel -> subtipo de retroalimentação.
PALAVRAS_RETRO = [
    ("poka", "Poka Yoke"), ("pfmea", "PFMEA"), ("trabalho padr", "Trabalho Padronizado"),
    ("plano de controle", "Plano de Controle"), ("sgi", "SGI (procedimentos)"), ("folha de prog", "Folha de programação"),
    ("aux", "Auxílio Visual Defeitos"), ("rea", "Plano Reação / Contig. / Emerg."), ("npr", "Revalidação do NPR"),
    ("lpa", "Layered Process Audit (LPA)"),
]


def subgrupo_retro(rotulo):
    rotulo = rotulo.lower()
    for palavra, sub in PALAVRAS_RETRO:
        if palavra in rotulo:
            return sub
    return None


def ler_a3(caminho):
    try:
        import openpyxl
    except ImportError as erro:  # pragma: no cover
        raise RuntimeError("Biblioteca openpyxl não instalada no servidor (pip install openpyxl).") from erro
    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = aba(wb, "FRP - A3")
    if ws is None:
        raise ValueError("aba 'FRP - A3' não encontrada (não é o formulário FM-QS-025?)")
    v = lambda c: texto(ws[c].value)  # noqa: E731

    numero = v("AV2").replace("\n", " ")
    partes = [p.strip() for p in re.split(r"[|/\s]+", numero) if p.strip()]
    id_mrhb = (partes[0] if partes else Path(caminho).stem.split(" - ")[0]).upper()
    notificacao = partes[1] if len(partes) > 1 else ""

    emissao = data_iso(ws["AD3"].value) or date.today().isoformat()
    prazo_limite = data_iso(ws["BC4"].value)
    analise = {
        "id_mrhb": id_mrhb,
        "numero_notificacao": notificacao or None,
        "responsavel": v("E3") or None,
        "area": v("O3") or None,
        "turno": v("V3") or None,
        "data_emissao": emissao,
        "area_gestor": v("AQ3") or None,
        "tema": v("E4") or "Resposta Rapida",
        "maquina": v("O4") or None,
        "supervisor": v("V4") or None,
        "gestor": v("AQ4") or None,
        "prazo_limite": prazo_limite or None,
        "cliente": v("B7") or None,
        "produto": v("I7") or None,
        "pn_cliente": v("Q7") or None,
        "pn_mrhb": v("Q8") or None,
        "quantidade": v("V7") or None,
        "definicao_problema": v("B23") or None,
    }

    dados = dados_vazios()
    descricao = v("B10")
    dados["descricao"]["descricao"] = re.sub(r"^descri[cç][aã]o:\s*", "", descricao, flags=re.I)
    dados["descricao"]["rastreabilidade_injecao"] = v("B21")
    dados["descricao"]["rastreabilidade_usinagem"] = v("H21")
    dados["descricao"]["objetivo"] = v("N23")

    for linha in range(27, 31):
        for col_nome, col_setor in (("B", "I"), ("M", "U")):
            if v(f"{col_nome}{linha}"):
                dados["equipe"].append({"nome": v(f"{col_nome}{linha}"), "setor": v(f"{col_setor}{linha}")})

    for chave, coluna in (("ocorrencia", "AA"), ("deteccao", "AI"), ("previsao", "AQ")):
        dados["porques"][chave]["respostas"] = [v(f"{coluna}{l}") for l in (9, 11, 13, 15, 17)]
        dados["porques"][chave]["causa_raiz"] = v(f"{coluna}19")

    ishikawa = {"material": "B", "mao_obra": "H", "maquina": "N", "medicao": "B", "metodo": "H", "meio_ambiente": "N"}
    for chave, coluna in ishikawa.items():
        linhas = range(46, 52) if chave in ("material", "mao_obra", "maquina") else range(56, 61)
        dados["ishikawa"][chave] = "\n".join(v(f"{coluna}{l}") for l in linhas if v(f"{coluna}{l}"))

    # Nº do alerta nas linhas fixas da 3.1 ("Elaborar Alerta da Qualidade nº E-70232026").
    for chave, celula in (("alerta_numero", "B35"), ("alerta_numero_contencao", "B36")):
        m = re.search(r"n[º°o]\s*(.+)$", v(celula))
        dados["contencao"][chave] = m.group(1).strip() if m else ""
    dados["contencao"]["ponto_corte_nf"] = v("J38")
    dados["contencao"]["exportacao_aplicavel"] = any(pessoa_valida(v(f"Q{l}")) for l in range(42, 45))
    dados["contencao"]["ponto_corte_qtd"] = v("N38")
    dados["fechamento"]["licoes_aprendidas"] = v("Z58")

    resultado = None
    for linha, rotulo in ((58, "8D Aprovado"), (59, "8D A/Condicional"), (60, "8D Reprovado")):
        if v(f"AW{linha}"):
            resultado = rotulo
    analise["resultado_validacao"] = resultado
    analise["status"] = "Em andamento"  # o status final é calculado ao gravar (ver gravar())

    acoes = []

    def acao(grupo, subgrupo, descricao, responsavel, prazo, farol, abrangencia=None, eficaz=None):
        responsavel = " ".join((responsavel or "").split())  # "Jonathas  Silva" -> "Jonathas Silva"
        if not descricao or descricao.rstrip(":").strip().lower() == "outros" or not pessoa_valida(responsavel):
            return
        # Na planilha, às vezes o "responsável" é um setor (ex.: "Qualidade", "SGI").
        setor = next((s for s in SETORES_PADRAO if s.lower() == responsavel.lower()), None)
        prazo = data_iso(prazo) or analise["prazo_limite"] or emissao
        status = STATUS_FAROL.get(texto(farol).upper(), "Aberta")
        acoes.append({
            "grupo": grupo, "subgrupo": subgrupo, "descricao": descricao, "responsavel": responsavel, "setor": setor,
            "data_abertura": emissao, "prazo": prazo, "status": status,
            "data_conclusao": prazo if status == "Concluída" else None,
            "abrangencia": abrangencia or None, "eficaz": eficaz or None,
        })

    for linha in range(34, 40):  # 3.1 Contenção nacional
        acao("contencao", "3.1 Nacional", v(f"B{linha}"), v(f"Q{linha}"), ws[f"T{linha}"].value, v(f"W{linha}"))
    for linha in range(42, 45):  # 3.2 Exportação
        acao("contencao", "3.2 Exportação", v(f"B{linha}"), v(f"Q{linha}"), ws[f"T{linha}"].value, v(f"W{linha}"))
    for linha in range(24, 33):  # 5 Corretivas
        tipo = v(f"Z{linha}").rstrip(":").strip()
        sub = next((s for s in ("Ocorrência", "Detecção", "Sistêmico") if s.lower()[:5] == tipo.lower()[:5]), None)
        acao("corretiva", sub, v(f"AB{linha}"), v(f"AP{linha}"), ws[f"AT{linha}"].value, v(f"AW{linha}"), v(f"AY{linha}"))
    for linha in range(34, 37):  # 5 Preventivas
        acao("preventiva", None, v(f"Z{linha}") or v(f"AB{linha}"), v(f"AP{linha}"), ws[f"AT{linha}"].value, v(f"AW{linha}"), v(f"AY{linha}"))
    for linha in range(41, 51):  # 6 Retroalimentação
        rotulo = v(f"Z{linha}")
        acao("retroalimentacao", subgrupo_retro(rotulo), v(f"AH{linha}") or rotulo.rstrip(":"), v(f"AP{linha}"),
             ws[f"AT{linha}"].value, v(f"AY{linha}"))
    for linha in range(53, 56):  # 7 Eficácia
        rotulo = v(f"Z{linha}")
        sub = ("Validação das ações implementadas" if rotulo.lower().startswith("valida")
               else "Análise de eficácia" if rotulo.lower().startswith("an") and "efic" in rotulo.lower()
               else "Auditoria Quality Foccus - SGI" if "auditoria" in rotulo.lower() else "Outra")
        eficaz = v(f"AY{linha}")
        acao("eficacia", sub, rotulo, v(f"AN{linha}"), ws[f"AQ{linha}"].value, v(f"AT{linha}"),
             eficaz=eficaz if eficaz in ("Eficaz", "Não eficaz") else None)

    # Check list
    wc = aba(wb, "Check Atividades Reclamação")
    if wc is not None:
        chaves = iter(k for k, _ in CHECKLIST)
        for linha in range(10, 44):
            if not texto(wc[f"A{linha}"].value):
                continue
            chave = next(chaves, None)
            if chave is None:
                break
            item = dados["checklist"]["itens"][chave]
            item["ok"] = bool(texto(wc[f"E{linha}"].value))
            item["na"] = bool(texto(wc[f"F{linha}"].value)) and not item["ok"]
            item["obs"] = texto(wc[f"G{linha}"].value)
        dados["checklist"]["observacoes"] = texto(wc["A46"].value)
        for chave, linha in (("atendimento", 51), ("gerente_qualidade", 53), ("sistema_qualidade", 55)):
            dados["checklist"]["assinaturas"][chave]["nome"] = texto(wc[f"D{linha}"].value)

    # Planilha de contenção
    wp = aba(wb, "Planilha de contenção")
    if wp is not None:
        locais = list(dados["contencao"]["locais"].keys())
        for i, linha in enumerate(range(9, 28)):
            if i < len(locais):
                registro = dados["contencao"]["locais"][locais[i]]
                for chave, col in (("suspeitas", "C"), ("aprovadas", "D"), ("reprovadas", "E"), ("re_executor", "F")):
                    registro[chave] = texto(wp[f"{col}{linha}"].value)

    # Check de causa raiz: blocos de 6 linhas por categoria (cabeçalho em 7, 15, 23, 31, 39, 47)
    wr = aba(wb, "Check de Causa Raiz")
    if wr is not None:
        for (chave, _), cabecalho in zip(CATEGORIAS_6M, (7, 15, 23, 31, 39, 47)):
            for linha in range(cabecalho + 1, cabecalho + 7):
                registro = {
                    "possivel_causa": texto(wr[f"B{linha}"].value), "especificado": texto(wr[f"K{linha}"].value),
                    "encontrado": texto(wr[f"Q{linha}"].value), "evidencias": texto(wr[f"W{linha}"].value), "confirmada": "",
                }
                if any(registro.values()):
                    dados["check_causa"][chave].append(registro)

    if not analise["prazo_limite"]:
        d = date.fromisoformat(emissao)
        analise["prazo_limite"] = date.fromordinal(d.toordinal() + DIAS_PRAZO_LIMITE_A3).isoformat()
    return analise, dados, acoes


def gravar(conn, analise, dados, acoes):
    if conn.execute("SELECT 1 FROM a3_analises WHERE id_mrhb = ?", (analise["id_mrhb"],)).fetchone():
        return None
    for a in acoes:
        a["setor"] = normalizar_setor(a.get("setor")) or None
    for membro in dados.get("equipe", []):
        membro["setor"] = normalizar_setor(membro.get("setor")) or ""
    colunas = list(analise.keys())
    cursor = conn.execute(
        f"INSERT INTO a3_analises ({', '.join(colunas)}, dados) VALUES ({', '.join('?' for _ in colunas)}, ?)",
        [*analise.values(), json.dumps(dados, ensure_ascii=False)],
    )
    analise_id = cursor.lastrowid
    for a in acoes:
        conn.execute(
            """INSERT INTO a3_acoes (analise_id, grupo, subgrupo, descricao, responsavel, setor, data_abertura, prazo,
                                     prazo_original, status, data_conclusao, abrangencia, eficaz)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (analise_id, a["grupo"], a["subgrupo"], a["descricao"], a["responsavel"], a["setor"], a["data_abertura"], a["prazo"],
             a["prazo"], a["status"], a["data_conclusao"], a["abrangencia"], a["eficaz"]),
        )
        conn.execute("INSERT OR IGNORE INTO a3_pessoas (nome, setor) VALUES (?, ?)", (a["responsavel"], a["setor"]))
    for membro in dados["equipe"]:
        membro["nome"] = " ".join(membro["nome"].split())
        conn.execute("INSERT OR IGNORE INTO a3_pessoas (nome, setor) VALUES (?, ?)", (membro["nome"], membro["setor"] or None))
        conn.execute("UPDATE a3_pessoas SET setor = COALESCE(setor, ?) WHERE nome = ?", (membro["setor"] or None, membro["nome"]))
    # Setor das ações = setor da pessoa na equipe (a planilha não tem setor por ação).
    conn.execute(
        """UPDATE a3_acoes SET setor = (SELECT setor FROM a3_pessoas p WHERE p.nome = a3_acoes.responsavel)
           WHERE analise_id = ? AND setor IS NULL""",
        (analise_id,),
    )
    # Mesma regra do sistema: Concluído = ações G + etapas do 8D realizadas + "8D Aprovado".
    status_automatico_a3(conn, analise_id)
    conn.execute(
        "INSERT INTO a3_historico (analise_id, id_mrhb, tipo, valor_novo) VALUES (?, ?, 'Importação do Excel', ?)",
        (analise_id, analise["id_mrhb"], f"{len(acoes)} ações importadas"),
    )
    return analise_id
